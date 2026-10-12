"""Synthetic boundary probes, complementing the small format fixtures."""

import io
import json
from pathlib import Path
import platform
import resource
import time
import zipfile

from openpyxl import Workbook, load_workbook
from PIL import Image
from pypdf import PdfWriter, PdfReader

from files import OUT, check, rejected, RESULTS


def validate_source_size(size):
    if size > 20*1024*1024:
        raise ValueError("20 MiB source limit")
    return True


def validate_workbook_metadata(archive):
    if sum(info.file_size for info in archive.infolist()) > 200*1024*1024:
        raise ValueError("200 MiB expanded workbook limit")


def scan_workbook(path):
    with zipfile.ZipFile(path) as archive:
        validate_workbook_metadata(archive)
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    if len(workbook.worksheets) > 50:
        workbook.close()
        raise ValueError("50-sheet limit")
    cells = 0
    try:
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if cell.value is not None:
                        if cell.row > 20_000 or cell.column > 200:
                            raise ValueError("20,000 rows / 200 populated columns")
                        cells += 1
                        if cells > 200_000:
                            raise ValueError("200,000 populated cells")
    finally:
        workbook.close()
    return cells


def raster_preview(raw):
    with Image.open(io.BytesIO(raw)) as image:
        width, height = image.size
        if width*height > 16_000_000 or max(width, height) > 8000:
            raise ValueError("16 million pixels / 8,000 px per side")
        image.load()
        preview = io.BytesIO()
        image.convert("RGB").save(preview, format="PNG")
        return preview.getvalue()


def pdf_pages(raw):
    reader = PdfReader(io.BytesIO(raw))
    if reader.is_encrypted or len(reader.pages) > 50:
        raise ValueError("Encrypted PDF or more than 50 pages")
    return len(reader.pages)


def run():
    OUT.mkdir(exist_ok=True, parents=True)
    started = time.perf_counter()
    check("upload: exact 20 MiB accepted", lambda: validate_source_size(20*1024*1024))
    check("upload: 20 MiB + 1 rejected", lambda: rejected(lambda: validate_source_size(20*1024*1024+1)), "ValueError")
    # Exercise the ZIP preflight independently without allocating a 200 MiB payload.
    class ArchiveMetadata:
        def infolist(self):
            info = zipfile.ZipInfo("xl/worksheets/sheet1.xml")
            info.file_size = 200*1024*1024+1
            return [info]
    check("xlsx: oversized expanded metadata rejected", lambda: rejected(lambda: validate_workbook_metadata(ArchiveMetadata())), "ValueError")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    for row in range(1000):
        sheet.append([row]*200)
    large = OUT / "boundary-200000.xlsx"
    workbook.save(large)
    check("xlsx: stream 200,000 cells", lambda: scan_workbook(large), 200_000)
    # Reuse the same source payload and append one populated cell to the worksheet XML.
    overflow = OUT / "boundary-200001.xlsx"
    with zipfile.ZipFile(large) as original, zipfile.ZipFile(overflow, "w", zipfile.ZIP_DEFLATED) as target:
        for info in original.infolist():
            raw = original.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                raw = raw.replace(b"</sheetData>", b'<row r="1001"><c r="A1001" t="n"><v>1</v></c></row></sheetData>')
            target.writestr(info.filename, raw)
    check("xlsx: 200,001 cells rejected", lambda: rejected(lambda: scan_workbook(overflow)), "ValueError")
    workbook = Workbook()
    for number in range(50):
        workbook.create_sheet(f"Sheet{number}")
    path = OUT / "too-many-sheets.xlsx"
    workbook.save(path)
    check("xlsx: 51 sheets rejected", lambda: rejected(lambda: scan_workbook(path)), "ValueError")
    workbook = Workbook()
    workbook.active["GR20000"] = 1
    path = OUT / "boundary-row-column.xlsx"
    workbook.save(path)
    check("xlsx: row 20,000 / column 200 accepted", lambda: scan_workbook(path), 1)
    for position in ["GS1", "A20001"]:
        workbook = Workbook()
        workbook.active[position] = 1
        path = OUT / f"too-wide-or-tall-{position}.xlsx"
        workbook.save(path)
        check(f"xlsx: populated {position} rejected", lambda path=path: rejected(lambda: scan_workbook(path)), "ValueError")
    for label, dimensions in [("long side", (8001,1)), ("pixel count", (4001,4000))]:
        buffer = io.BytesIO()
        Image.new("1", dimensions).save(buffer, format="PNG")
        check(f"raster: {label} rejected before load", lambda buffer=buffer: rejected(lambda: raster_preview(buffer.getvalue())), "ValueError")
    check("raster: malformed rejected", lambda: rejected(lambda: raster_preview(b"not an image")), "UnidentifiedImageError")
    check("raster: truncated rejected", lambda: rejected(lambda: raster_preview((OUT / "synthetic.png").read_bytes()[:100])), "OSError")
    writer = PdfWriter()
    for _ in range(50):
        writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    check("pdf: 50 pages accepted", lambda: pdf_pages(buffer.getvalue()), 50)
    writer.add_blank_page(width=100,height=100)
    overflow = io.BytesIO()
    writer.write(overflow)
    check("pdf: 51 pages rejected", lambda: rejected(lambda: pdf_pages(overflow.getvalue())), "ValueError")
    check("pdf: encrypted rejected", lambda: rejected(lambda: pdf_pages((OUT / "encrypted.pdf").read_bytes())), "ValueError")
    rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    if platform.system() != "Darwin":
        rss *= 1024
    report = {"total_seconds_including_fixture_generation": time.perf_counter()-started, "peak_rss_mib": rss/1024**2,
              "valid_200000_cell_source_bytes": large.stat().st_size, "cases": RESULTS}
    (OUT / "limits-results.json").write_text(json.dumps(report, indent=2)+"\n")
    print(json.dumps({"passed": sum(r["passed"] for r in RESULTS), "total": len(RESULTS), "failures": [r for r in RESULTS if not r["passed"]],
                      "peak_rss_mib": report["peak_rss_mib"], "seconds": report["total_seconds_including_fixture_generation"]}, indent=2))
    return all(r["passed"] for r in RESULTS)


if __name__ == "__main__":
    raise SystemExit(0 if run() else 1)

"""Disposable M0 format probes, using generated public-safe fixtures only.

These checks establish feasibility; they are not the production upload adapters.
"""

import base64
import csv
from datetime import datetime
import hashlib
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import io
import json
import os
from pathlib import Path
import re
import subprocess
import tarfile
import threading
import time
from urllib.parse import quote, unquote
import urllib.request
import zipfile
import zlib

from defusedxml import ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from pypdf.actions import JavaScript
from pypdf.generic import ArrayObject, DecodedStreamObject, DictionaryObject, NameObject, FloatObject, TextStringObject
import cairosvg

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / ".qualification"
OUT = WORK / "preview"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
RESULTS = []


def check(name, operation, expected=None):
    start = time.perf_counter()
    try:
        value = operation()
        ok = value == expected if expected is not None else bool(value)
        detail = value
    except Exception as exc:
        ok, detail = False, f"{type(exc).__name__}: {exc}"
    RESULTS.append({"case": name, "passed": bool(ok), "seconds": round(time.perf_counter()-start, 4), "detail": detail})


def rejected(operation):
    try:
        operation()
    except Exception as exc:
        # Exception class is recorded; a probe rejection is not a production error contract.
        return type(exc).__name__
    raise AssertionError("Unsafe/invalid fixture was accepted")


def compress(xml):
    encoder = zlib.compressobj(wbits=-15)
    return base64.b64encode(encoder.compress(quote(xml, safe="~()*!.'").encode()) + encoder.flush()).decode()


def bounded_xml(raw, limit=8*1024*1024):
    if len(raw) > limit or b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        raise ValueError("XML size/DTD/entity policy")
    return ET.fromstring(raw)


def drawio_pages(raw):
    root = bounded_xml(raw)
    pages = root.findall("diagram") if root.tag == "mxfile" else [root]
    if len(pages) > 50:
        raise ValueError("50-page limit")
    output = []
    for page in pages:
        model = page.find("mxGraphModel") if page.tag == "diagram" else page
        if model is None:
            decoder = zlib.decompressobj(-15)
            decoded = decoder.decompress(base64.b64decode(page.text, validate=True), 8*1024*1024+1)
            if len(decoded) > 8*1024*1024 or not decoder.eof:
                raise ValueError("Expanded draw.io page limit")
            model = bounded_xml(unquote(decoded.decode()).encode())
        if model.tag != "mxGraphModel":
            raise ValueError("Expected mxGraphModel")
        # This conservative preflight is deliberately narrower than the future adapter.
        for node in model.iter():
            for key, value in node.attrib.items():
                if re.search(r"https?://|file:|javascript:|<script|<iframe|onerror\s*=", value, re.I):
                    raise ValueError(f"External/active draw.io attribute: {key}")
        output.append(ET.tostring(model, encoding="unicode"))
    return output


def workbook_probe():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Données fictives"
    for row in [["Application", "Montant", "Date", "Taux"], ["DEMO-A", 1250.5, datetime(2026, 10, 11), .125],
                ["Formule enregistrée", "=1+2"], ["Formule vide", '=IF(1=1,"","x")'], ["Sans résultat", "=7*8"], ["Cache numérique vide", "=8*9"]]:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet["B2"].number_format = "#,##0.00"
    sheet["C2"].number_format = "dd/mm/yyyy"
    sheet["D2"].number_format = "0.0%"
    workbook.create_sheet("Masquée").sheet_state = "hidden"
    workbook["Masquée"]["A1"] = "Feuille masquée — contenu fictif"
    buffer = io.BytesIO()
    workbook.save(buffer)
    # openpyxl does not calculate formulas. Construct three distinct cache states.
    target = OUT / "synthetic.xlsx"
    with zipfile.ZipFile(buffer) as original, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as patched:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                xml = bounded_xml(data)
                cells = {c.attrib["r"]: c for c in xml.findall(".//s:c", NS)}
                cells["B3"].find("s:v", NS).text = "3"
                cells["B4"].set("t", "str")
                cells["B5"].remove(cells["B5"].find("s:v", NS))
                data = ET.tostring(xml)
            patched.writestr(info.filename, data)
    with zipfile.ZipFile(target) as archive:
        if sum(i.file_size for i in archive.infolist()) > 200*1024*1024:
            raise ValueError("Expanded workbook limit")
        raw_sheet = bounded_xml(archive.read("xl/worksheets/sheet1.xml"))
    caches = {}
    for cell in raw_sheet.findall(".//s:c", NS):
        if cell.find("s:f", NS) is not None:
            value = cell.find("s:v", NS)
            # An empty numeric <v/> is commonly an unsaved result, not an empty string.
            caches[cell.attrib["r"]] = "missing" if value is None else "empty" if value.text is None and cell.get("t") == "str" else "missing" if value.text is None else value.text
    values = load_workbook(target, data_only=True, keep_links=False)
    formulas = load_workbook(target, data_only=False, keep_links=False)
    manifest = []
    for sheet in values:
        rows = []
        for row in sheet:
            cells = []
            for cell in row:
                value = cell.value
                if caches.get(cell.coordinate) == "missing" and sheet == values.worksheets[0]:
                    value = "Résultat non disponible"
                elif isinstance(value, datetime):
                    value = value.strftime("%d/%m/%Y")
                elif isinstance(value, (int, float)) and "%" in cell.number_format:
                    value = f"{value*100:.1f}%"
                elif isinstance(value, (int, float)) and "0.00" in cell.number_format:
                    value = f"{value:,.2f}"
                cells.append({"value": "" if value is None else str(value), "bold": cell.font.bold})
            rows.append(cells)
        manifest.append({"name": sheet.title, "hidden": sheet.sheet_state != "visible", "rows": rows})
    (OUT / "workbook.json").write_text(json.dumps(manifest, ensure_ascii=False))
    assert caches == {"B3": "3", "B4": "empty", "B5": "missing", "B6": "missing"}, caches
    assert values.worksheets[0]["B3"].value == 3
    assert values.worksheets[0]["B4"].value is None
    assert formulas.worksheets[0]["B5"].value == "=7*8"
    return {"cache_states": caches, "sheets": len(manifest), "hidden_included": manifest[1]["hidden"]}


def parse_csv(raw, delimiter=None):
    text = raw.decode("utf-8-sig")
    candidates = {}
    for candidate in ([delimiter] if delimiter else [",", ";", "\t"]):
        rows = list(csv.reader(io.StringIO(text), delimiter=candidate, strict=True))
        if rows and len(rows[0]) > 1 and all(len(row) == len(rows[0]) for row in rows):
            candidates[candidate] = rows
    if len(candidates) != 1:
        raise ValueError("Choose a delimiter or correct inconsistent rows")
    return next(iter(candidates.values()))


def safe_svg(raw):
    root = bounded_xml(raw)
    allowed = {"svg", "g", "rect", "circle", "ellipse", "line", "path", "polyline", "polygon", "text", "tspan", "defs", "marker", "title", "desc"}
    for node in root.iter():
        if node.tag.split("}")[-1] not in allowed:
            raise ValueError("Unsupported/active SVG element")
        for key, value in node.attrib.items():
            key = key.split("}")[-1]
            if key.startswith("on") or key in {"href", "src", "style"} or re.search(r"https?:|file:|url\(", value, re.I):
                raise ValueError("Unsupported/active SVG attribute")
    return cairosvg.svg2png(bytestring=ET.tostring(root), output_width=720, output_height=220)


def plantuml(source, name, timeout=20, suffix=".puml", profile="SANDBOX"):
    target = OUT / f"{name}{suffix}"
    target.write_text(source)
    environment = {**os.environ, "PLANTUML_SECURITY_PROFILE": profile}
    completed = subprocess.run([
        "java", "-Djava.awt.headless=true", f"-DPLANTUML_SECURITY_PROFILE={profile}", "-Xmx256m", "-jar", str(WORK / "assets/plantuml.jar"),
        "-Playout=smetana", "-tpng", "-failfast2", str(target),
    ], env=environment, capture_output=True, text=True, timeout=timeout)
    images = sorted(OUT.glob(f"{name}*.png"))
    return completed.returncode, images, completed.stderr[-1000:]


def plantuml_remote_probe():
    requests = []

    class IncludeHandler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path)
            body = b"class M0ReachableSyntheticInclude\n"
            self.send_response(200)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), IncludeHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        url = f"http://127.0.0.1:{server.server_port}/synthetic-include.puml"
        with urllib.request.urlopen(url, timeout=5) as response:
            assert response.status == 200
        source = f"@startuml\n!include {url}\n@enduml"
        # Trusted positive control, restricted to this synthetic loopback resource.
        requests.clear()
        control_code, _, control_error = plantuml(source, "remote-control", profile="UNSECURE")
        assert control_code == 0 and requests, f"Remote control: exit={control_code}, requests={len(requests)}, diagnostic={control_error}"
        control_access_observed = len(requests) > 0
        requests.clear()
        code, _, error = plantuml(source, "denied-remote")
        return {"control_succeeded": control_code == 0, "control_access_observed": control_access_observed, "sandbox_refused": code != 0 and bool(error), "sandbox_requests": len(requests)}
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def pdf_fixture():
    writer = PdfWriter()
    for index in range(2):
        page = writer.add_blank_page(width=520, height=220)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 24 Tf 30 150 Td (Diagramme fictif - page {index+1}) Tj ET 30 45 200 70 re S 280 45 180 70 re S 230 80 m 280 80 l S".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    target = OUT / "synthetic.pdf"
    writer.add_open_action(JavaScript("app.alert('Synthetic M0: scripts must never execute');"))
    annotation = DictionaryObject({NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Link"),
        NameObject("/Rect"): ArrayObject([FloatObject(x) for x in [30, 45, 230, 115]]),
        NameObject("/A"): DictionaryObject({NameObject("/S"): NameObject("/URI"), NameObject("/URI"): TextStringObject("https://example.invalid/m0")})})
    writer.pages[0][NameObject("/Annots")] = ArrayObject([writer._add_object(annotation)])
    writer.write(target)
    assert len(PdfReader(target).pages) == 2
    assert len(PdfReader(target).trailer["/Root"]["/Names"]["/JavaScript"]["/Names"]) == 2
    writer.encrypt("synthetic-only")
    writer.write(OUT / "encrypted.pdf")
    assert PdfReader(OUT / "encrypted.pdf").is_encrypted
    return {"pages": 2, "encrypted_detected": True, "script_and_external_link_present_for_browser_check": True}


def generate():
    OUT.mkdir(parents=True, exist_ok=True)
    check("xlsx: cached numeric/empty/missing, styles, hidden sheet", workbook_probe)
    raw_csv = '\ufeffApplication;Description\nDEMO-A;"Première ligne\nDeuxième ligne"\nDEMO-B;"Avec ; séparateur"\n'.encode()
    (OUT / "synthetic.csv").write_bytes(raw_csv)
    csv_rows = parse_csv(raw_csv)
    (OUT / "csv.json").write_text(json.dumps({"name": "CSV fictif", "rows": [[{"value": value, "bold": index == 0} for value in row] for index, row in enumerate(csv_rows)]}, ensure_ascii=False))
    check("csv: BOM, semicolon, quoted newline", lambda: len(parse_csv(raw_csv)), 3)
    check("csv: comma", lambda: parse_csv(b"a,b\n1,2\n"), [["a", "b"], ["1", "2"]])
    check("csv: tab", lambda: parse_csv(b"a\tb\n1\t2\n"), [["a", "b"], ["1", "2"]])
    check("csv: ambiguous rejected", lambda: rejected(lambda: parse_csv(b"a,b;c\n1,2;3\n")), "ValueError")
    check("csv: ambiguous override", lambda: parse_csv(b"a,b;c\n1,2;3\n", ";"), [["a,b", "c"], ["1,2", "3"]])
    check("csv: inconsistent rejected", lambda: rejected(lambda: parse_csv(b"a;b\n1;2;3\n")), "ValueError")
    check("csv: bad encoding rejected", lambda: rejected(lambda: parse_csv(b"a;b\n\xff;2\n")), "UnicodeDecodeError")
    image = Image.new("RGB", (520, 160), "#f1f5f9")
    draw = ImageDraw.Draw(image)
    draw.rectangle((25, 35, 210, 120), outline="#164e63", width=3)
    draw.rectangle((300, 35, 490, 120), outline="#164e63", width=3)
    draw.line((210, 80, 300, 80), fill="#164e63", width=3)
    draw.text((50, 70), "Application DEMO-A", fill="black")
    draw.text((320, 70), "Application DEMO-B", fill="black")
    for extension in ["png", "jpeg"]:
        image.save(OUT / f"synthetic.{extension}")
        check(f"{extension}: decode/re-encode", lambda ext=extension: Image.open(OUT / f"synthetic.{ext}").size, (520, 160))
    (OUT / "synthetic.jpg").write_bytes((OUT / "synthetic.jpeg").read_bytes())
    check("jpg: JPEG extension alias", lambda: Image.open(OUT / "synthetic.jpg").format, "JPEG")
    svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="720" height="220"><rect x="25" y="40" width="230" height="110" fill="#e0f2fe" stroke="#164e63"/><text x="50" y="100" font-size="22">Base DEMO</text><path d="M255 95 L440 95" stroke="#164e63"/><rect x="440" y="40" width="230" height="110" fill="#dcfce7" stroke="#166534"/><text x="465" y="100" font-size="22">Service fictif</text></svg>'
    (OUT / "synthetic.svg").write_bytes(svg)
    (OUT / "svg-preview.png").write_bytes(safe_svg(svg))
    check("svg: local rasterization", lambda: Image.open(OUT / "svg-preview.png").size, (720, 220))
    for label, bad in [("script", b'<svg><script>alert(1)</script></svg>'), ("external image", b'<svg><image href="https://example.invalid/x"/></svg>'), ("entity", b'<!DOCTYPE svg [<!ENTITY x SYSTEM "file:///synthetic">]><svg>&x;</svg>')]:
        check(f"svg: {label} rejected", lambda bad=bad: rejected(lambda: safe_svg(bad)), "ValueError")
    check("pdf: two pages and encrypted detection", pdf_fixture)
    source = '@startuml\nentity Customer {\n * id : UUID\n name : text\n}\nentity Booking {\n * id : UUID\n customer_id : UUID\n}\nCustomer ||--o{ Booking\n@enduml\n@startuml\nclass DemoApplication\nclass DemoDatabase\nDemoApplication --> DemoDatabase\n@enduml\n'
    code, images, error = plantuml(source, "database")
    check("plantuml: SANDBOX, Smetana, two blocks", lambda: {"exit": code, "pages": len(images)}, {"exit": 0, "pages": 2})
    alias_code, alias_images, _ = plantuml("@startuml\nclass Demo\n@enduml", "alias", suffix=".plantuml")
    check("plantuml: .plantuml extension alias", lambda: alias_code == 0 and len(alias_images) == 1)
    (OUT / "local-sentinel.txt").write_text("class M0_PRIVATE_SENTINEL")
    for label, include in [("local", str(OUT / "local-sentinel.txt"))]:
        code, images, error = plantuml(f"@startuml\n!include {include}\n@enduml", f"denied-{label}")
        check(f"plantuml: {label} include blocked", lambda code=code, error=error: {"nonzero": code != 0, "diagnostic": bool(error)}, {"nonzero": True, "diagnostic": True})
    check("plantuml: reachable remote include blocked (positive control)", plantuml_remote_probe,
          {"control_succeeded": True, "control_access_observed": True, "sandbox_refused": True, "sandbox_requests": 0})
    code, _, _ = plantuml("@startuml\nthis is not valid syntax !@#\n@enduml", "invalid")
    check("plantuml: invalid syntax rejected", lambda: code != 0)
    check("plantuml: subprocess timeout terminates renderer", lambda: rejected(lambda: plantuml(source, "timeout", timeout=.001)), "TimeoutExpired")
    picture = base64.b64encode((OUT / "synthetic.png").read_bytes()).decode()
    stencil = compress('<shape name="demo" w="100" h="60" aspect="variable" strokewidth="inherit"><background><rect x="0" y="0" w="100" h="60"/></background><foreground><fillstroke/><path><move x="0" y="0"/><line x="100" y="60"/></path><stroke/></foreground></shape>')
    model = f'<mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/><mxCell id="2" value="&lt;b&gt;Application DEMO-A&lt;/b&gt;" style="rounded=1;html=1;fillColor=#e0f2fe;" vertex="1" parent="1"><mxGeometry x="30" y="30" width="190" height="70" as="geometry"/></mxCell><mxCell id="3" value="Base fictive" style="shape=cylinder;fillColor=#dcfce7;" vertex="1" parent="1"><mxGeometry x="300" y="30" width="150" height="70" as="geometry"/></mxCell><mxCell id="4" edge="1" parent="1" source="2" target="3"><mxGeometry relative="1" as="geometry"/></mxCell><mxCell id="5" value="Stencil intégré" style="shape=stencil({stencil});" vertex="1" parent="1"><mxGeometry x="35" y="150" width="150" height="70" as="geometry"/></mxCell><mxCell id="6" style="shape=image;image=data:image/png,{picture};" vertex="1" parent="1"><mxGeometry x="250" y="140" width="260" height="80" as="geometry"/></mxCell></root></mxGraphModel>'
    # A named shape exercises the separately vendored stencil bundle as well.
    model = model.replace("shape=cylinder", "shape=mxgraph.flowchart.database")
    second = model.replace("Application DEMO-A", "Page compressée DEMO-B")
    mxfile = f'<mxfile><diagram name="Flux principal">{model}</diagram><diagram name="Page compressée">{compress(second)}</diagram></mxfile>'
    (OUT / "synthetic.drawio").write_text(mxfile)
    (OUT / "synthetic.xml").write_text(model)
    check("drawio: compressed and uncompressed pages", lambda: len(drawio_pages(mxfile.encode())), 2)
    check("drawio: standalone XML", lambda: len(drawio_pages(model.encode())), 1)
    check("drawio: external image rejected", lambda: rejected(lambda: drawio_pages(model.replace(f"data:image/png,{picture}", "https://example.invalid/image.png").encode())), "ValueError")
    check("drawio: DTD rejected", lambda: rejected(lambda: drawio_pages(b'<!DOCTYPE mxfile><mxfile/>')), "ValueError")
    check("drawio: 51 pages rejected", lambda: rejected(lambda: drawio_pages(('<mxfile>'+f'<diagram>{model}</diagram>'*51+'</mxfile>').encode())), "ValueError")
    check("drawio: 50 pages accepted", lambda: len(drawio_pages(('<mxfile>'+f'<diagram>{model}</diagram>'*50+'</mxfile>').encode())), 50)
    check("drawio: decompression bound", lambda: rejected(lambda: drawio_pages(f'<mxfile><diagram>{compress("x"*(8*1024*1024+1))}</diagram></mxfile>'.encode())), "ValueError")
    pages = drawio_pages(mxfile.encode())
    (OUT / "drawio-pages.json").write_text(json.dumps(pages))
    for target in [OUT / name for name in ["synthetic.xlsx", "synthetic.csv", "synthetic.drawio", "synthetic.xml", "database.puml", "alias.plantuml", "synthetic.pdf", "synthetic.png", "synthetic.jpeg", "synthetic.jpg", "synthetic.svg"]]:
        before = hashlib.sha256(target.read_bytes()).hexdigest()
        destination = OUT / "downloads" / target.name
        destination.parent.mkdir(exist_ok=True)
        destination.write_bytes(target.read_bytes())
        check(f"original bytes: {target.suffix}", lambda destination=destination, before=before: hashlib.sha256(destination.read_bytes()).hexdigest() == before)
    # Extract only the two needed, regular PDF.js build files; never extract arbitrary paths.
    with tarfile.open(WORK / "assets/pdfjs.tgz") as archive:
        for name in ["pdf.mjs", "pdf.worker.mjs"]:
            (OUT / name).write_bytes(archive.extractfile(f"package/build/{name}").read())
    (OUT / "files-results.json").write_text(json.dumps(RESULTS, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps({"passed": sum(x["passed"] for x in RESULTS), "total": len(RESULTS), "failures": [x for x in RESULTS if not x["passed"]]}, ensure_ascii=False, indent=2))
    return all(x["passed"] for x in RESULTS)


if __name__ == "__main__":
    raise SystemExit(0 if generate() else 1)

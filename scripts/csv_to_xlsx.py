#!/usr/bin/env python3
"""Convert every data/*.csv into Excel files.

Writes to the output directory (default: dist/):
  - one <name>.xlsx per CSV
  - UBSI_referentiels.xlsx with one sheet per CSV

Accepts CSV saved by Excel (";" separator, UTF-8 with BOM) as well as "," CSV.
Exits with code 1 if a CSV cannot be read.
"""
from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
NUMBER = re.compile(r"^-?(0|[1-9]\d*)([.,]\d+)?$")  # no leading zeros: "007" stays text


def read_csv(path: Path) -> list[list[str]]:
    text = path.read_text(encoding="utf-8-sig")
    first_line = text.splitlines()[0] if text.strip() else ""
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    rows = list(csv.reader(text.splitlines(), delimiter=delimiter))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise ValueError("empty file")
    width = len(rows[0])
    for i, r in enumerate(rows[1:], start=2):
        if len(r) > width:
            raise ValueError(f"line {i} has {len(r)} cells, header has {width}")
    return rows


def typed(value: str):
    v = value.strip()
    if NUMBER.match(v):
        n = float(v.replace(",", "."))
        return int(n) if n.is_integer() and "." not in v and "," not in v else n
    return v


def fill_sheet(ws, rows: list[list[str]]) -> None:
    ws.append(rows[0])
    for r in rows[1:]:
        ws.append([typed(c) for c in r])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(zip(*[r + [""] * (len(rows[0]) - len(r)) for r in rows]), start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, *(len(str(c)) + 2 for c in col)))


def main(out_dir: Path) -> int:
    out_dir.mkdir(parents=True, exist_ok=True)
    csvs = sorted(DATA_DIR.glob("*.csv"))
    combined = Workbook()
    combined.remove(combined.active)
    errors = 0
    for path in csvs:
        try:
            rows = read_csv(path)
        except (ValueError, UnicodeDecodeError, csv.Error) as e:
            print(f"ERROR data/{path.name}: {e}")
            errors += 1
            continue
        single = Workbook()
        single.active.title = path.stem[:31]
        fill_sheet(single.active, rows)
        single.save(out_dir / f"{path.stem}.xlsx")
        fill_sheet(combined.create_sheet(path.stem[:31]), rows)
        print(f"OK   data/{path.name} -> {path.stem}.xlsx ({len(rows) - 1} rows)")
    if combined.sheetnames:
        combined.save(out_dir / "UBSI_referentiels.xlsx")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main(Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "dist"))
